import io
import json
from datetime import datetime
from unittest.mock import MagicMock, patch

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from rest_framework.test import APIClient

from authentication.models import SystemConfig, UserProfile
from .models import Candidate, CandidateParticipation, CandidateSheetOutbox, ExamRoom, ExamSession, ExaminationBillingRecord, ExaminationSheet, RoundResult, SchoolRegistration, SessionSheetOutbox, TransferProof
from .proof_archive import archive_pending, folder_for
from .school_import import read_workbook


class SchoolImportTests(TestCase):
    def setUp(self):
        # Seed migrations contain real seasons; isolate the school fixture.
        self.client = APIClient()
        self.user = UserProfile.objects.create(email='exam-import@example.test', role='EMPLOYEE', access_modules=['examination'])
        self.finance = UserProfile.objects.create(email='finance-import@example.test', role='ADMIN')
        self.client.force_authenticate(self.user)
        self.sessions = {}
        for code in ('TESTA', 'TESTB'):
            self.sessions[code] = ExamSession.objects.create(id=f'{code}-2026', competition_id=code, code=code, name=code, parent=code, organizer='Test', time='2026–2027', sort_key=code,
                rounds=[{'id': 'round-1', 'name': 'Vòng 1', 'date': '2026-10-25'}])
            ExaminationSheet.objects.create(id=f'sheet-{code}', name=code, session_id=f'{code}-2026', url='https://docs.google.com/spreadsheets/d/test', sheet_tab=code, stage='session-output', created_at=timezone.now(), updated_at=timezone.now())
        self.options = {'partner': {'school': 'Trường A', 'representative': 'Người liên lạc', 'phone': '0912345678', 'email': 'school@example.test'}, 'academicYear': '2026-2027'}
        self.rows = [
            [1, 'Nguyễn Minh An', '12/07/2015', '001215012345', 'Lớp 6', 'TESTA', 250000, 'Nguyễn Văn A', '0901234567', 'parent@example.test', ''],
            [2, 'Nguyễn Minh An', '12/07/2015', '001215012345', 'Lớp 6', 'TESTB', 450000, 'Nguyễn Văn A', '0901234567', 'parent@example.test', ''],
            [3, 'Trần Minh Bình', '10/02/2014', '002214012345', 'Lớp 7', 'TESTA', 250000, 'Trần Văn A', '0907654321', 'parent2@example.test', ''],
        ]

    def workbook(self, rows=None, metadata=False, copy=False, school_headers=False):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Đăng ký tham dự'
        if metadata:
            sheet.append(['Tên trường / đơn vị (*):', None, 'Trường A'])
            sheet.append(['Người phụ trách (*):', 'Người liên lạc'])
            sheet.append(['Số điện thoại (*):', '0912345678'])
            sheet.append(['Email nhận thông tin (*):', None, 'school@example.test'])
        if school_headers:
            sheet.append(['STT', 'Họ và tên thí sinh', 'Ngày, tháng,\nnăm sinh', 'Căn cước công dân/\nHộ chiếu', 'Lớp đang\nhọc', 'Cuộc thi đăng ký\n(FIMO/FIEO)', 'Họ và tên phụ huynh/\nngười giám hộ', 'Số điện thoại', 'Email học sinh/\nphụ huynh', 'Ghi chú'])
        else:
            sheet.append(['STT', 'Họ và tên thí sinh (*)', 'Ngày sinh (*)\n(DD/MM/YYYY)', 'CCCD / Hộ chiếu / SĐD Học sinh', 'Lớp đang học (*)', 'Cuộc thi đăng ký (*)\n(Chọn danh sách)', 'Lệ phí dự thi\n(Tự động tính)', 'Họ tên phụ huynh / Người giám hộ (*)', 'Số điện thoại (*)\n(Nhận Zalo/SMS)', 'Email liên hệ (*)', 'Ghi chú'])
        for row in self.rows if rows is None else rows:
            sheet.append(row[:6] + row[7:] if school_headers else row)
        if copy:
            workbook.copy_worksheet(sheet).title = 'Bản sao danh sách'
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def post(self, content, options):
        return self.client.post('/api/examination/import/school', {'file': SimpleUploadedFile('school.xlsx', content), 'options': json.dumps(options)}, format='multipart')

    def preview(self, content=None, options=None):
        content = content or self.workbook()
        response = self.post(content, options or self.options)
        self.assertEqual(response.status_code, 200, response.data)
        return response

    def commit(self, content=None, options=None):
        content = content or self.workbook()
        options = options or self.options
        preview = self.preview(content, options)
        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        with patch('examination.partner_contact_sync.launch_partner_contact_sync'):
            response = self.post(content, options | {'action': 'commit', 'previewToken': preview.data['previewToken']})
        self.assertEqual(response.status_code, 201, response.data)
        return response

    def test_import_multiple_subjects_once_groups_school_and_queues_sheets(self):
        preview = self.preview()
        self.assertEqual(preview.data['summary']['candidates'], 2)
        self.assertEqual(Candidate.objects.count(), 0)
        self.assertFalse(SystemConfig.objects.filter(key='examination_partners').exists())
        result = self.commit()
        self.assertEqual(result.data['summary']['newCandidates'], 2)
        self.assertEqual(result.data['summary']['newRegistrations'], 3)
        report = {item['competitionCode']: item for item in result.data['report']}
        self.assertEqual(report['TESTA']['registrations'], 2)
        self.assertEqual(report['TESTB']['registrations'], 1)
        self.assertEqual(Candidate.objects.count(), 2)
        self.assertEqual(CandidateParticipation.objects.count(), 3)
        self.assertEqual(SchoolRegistration.objects.count(), 2)
        self.assertEqual(ExaminationBillingRecord.objects.count(), 2)
        self.assertEqual(CandidateSheetOutbox.objects.count(), 2)
        self.assertEqual(SessionSheetOutbox.objects.count(), 3)
        self.client.force_authenticate(self.finance)
        records = self.client.get('/api/examination/billing/records?scope=all')
        self.assertEqual(len(records.data), 2)
        self.assertEqual({r['kind'] for r in records.data}, {'school'})
        self.assertEqual(sum(r['amount'] for r in records.data), 950000)
        self.assertEqual(self.client.get('/api/examination/billing/stats?scope=all').data['totalAmount'], 950000)

    def test_school_headers_read_identity_parent_and_excel_dates(self):
        row = self.rows[0][:]
        row[2] = datetime(2015, 7, 12)
        row[3] = 1215012345
        content = self.workbook([row], school_headers=True)

        raw, _, _, _ = read_workbook(content)

        self.assertEqual(raw[0]['cccd'], '1215012345')
        self.assertEqual(raw[0]['dob'], '2015-07-12')
        self.assertEqual(raw[0]['parent'], 'Nguyễn Văn A')
        self.assertEqual(raw[0]['phone'], '0901234567')
        self.commit(content)
        candidate = Candidate.objects.get()
        self.assertEqual(candidate.identity, '001215012345')
        self.assertEqual(candidate.birth_date, '2015-07-12')

    def test_combined_codes_create_both_registrations_and_retry_is_idempotent(self):
        row = self.rows[0][:]
        for separator in (' & ', '/', ', ', '; ', ' + ', ' và ', '\n'):
            with self.subTest(separator=separator):
                row[5] = f'TESTA{separator}TESTB{separator}TESTA'
                content = self.workbook([row], school_headers=True)
                preview = self.preview(content)
                self.assertEqual(preview.data['summary']['rows'], 1)
                self.assertEqual(preview.data['summary']['candidates'], 1)
                self.assertEqual(preview.data['summary']['registrations'], 2)
                self.assertEqual({route['contest'] for route in preview.data['routes']}, {'TESTA', 'TESTB'})
                self.assertEqual({r['sessionId'] for r in preview.data['rows']}, {s.pk for s in self.sessions.values()})
                self.commit(content)
        self.assertEqual(Candidate.objects.count(), 1)
        self.assertEqual(CandidateParticipation.objects.count(), 2)
        self.assertEqual(SchoolRegistration.objects.count(), 2)
        self.assertEqual(ExaminationBillingRecord.objects.count(), 2)
        self.assertTrue(all(amount is None for amount in ExaminationBillingRecord.objects.values_list('amount', flat=True)))

    def test_combined_contests_split_one_fee_instead_of_double_billing(self):
        row = self.rows[0][:]
        row[5] = 'TESTA & TESTB'
        row[6] = '500.000VNĐ'
        preview = self.preview(self.workbook([row]))
        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        self.assertEqual({group['amount'] for group in preview.data['groups']}, {250000})

        row[6] = 250001
        content = self.workbook([row])
        preview = self.preview(content)
        self.assertFalse(preview.data['canCommit'])
        self.assertTrue(any('phí chung' in issue['message'] for issue in preview.data['issues']))
        response = self.post(content, self.options | {'action': 'commit', 'previewToken': preview.data['previewToken']})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Candidate.objects.count(), 0)
        self.assertEqual(ExaminationBillingRecord.objects.count(), 0)

    def appendix_workbook(self, rows):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Cả trường'
        sheet.append(['Tên trường học/tổ chức', None, 'Trường A'])
        sheet.append(['Người phụ trách', None, 'Họ và tên: Người Liên Lạc. Chức vụ: Hiệu phó.'])
        sheet.append(['Thông tin liên hệ', None, 'Số điện thoại: 0912345678.', None, 'Email: school@example.test'])
        sheet.append(['STT', 'Họ và tên thí sinh', 'Ngày, tháng,\nnăm sinh', 'Căn cước công dân/\nHộ chiếu', 'Lớp đang\nhọc', 'Cuộc thi đăng ký\n(FIMO/FIEO)', 'Họ và tên phụ huynh/\nngười giám hộ', 'Số điện thoại', 'Email học sinh/\nphụ huynh', 'Lệ phí'])
        for row in rows:
            sheet.append(row)
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()

    def test_reconcile_command_repairs_swapped_birth_dates_and_adds_missing_registrations(self):
        import tempfile
        from django.core.management import call_command
        old = Candidate.objects.create(id='FT-00562', code='FT-00562', name='Nguyễn Bảo Hoàng', birth_date='2015-10-06',
            identity='001215047377', class_name='6A1', sort_key='old')
        CandidateParticipation.objects.create(candidate=old, session=self.sessions['TESTA'])
        content = self.appendix_workbook([
            [1, 'Nguyễn Bảo Hoàng', datetime(2015, 6, 10), '001215047377', '6A1', 'TESTA & TESTB', 'Mai Thu Trang', '0984127270', 'mtt@example.test', '500.000VNĐ'],
            [2, 'Trần Minh Bình', datetime(2015, 2, 14), '1315028059', '6A2', 'TESTB', 'Trần Văn A', '907654321', 'b @example.test', '250.000VNĐ'],
        ])
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as handle:
            handle.write(content)
        arguments = ['reconcile_school_registrations', '--file', handle.name, '--academic-year', '2026-2027']
        with patch('examination.partner_contact_sync.launch_partner_contact_sync'):
            call_command(*arguments, stdout=io.StringIO())
            old.refresh_from_db()
            self.assertEqual(old.birth_date, '2015-10-06')
            self.assertEqual(Candidate.objects.count(), 1)
            output = io.StringIO()
            call_command(*arguments, '--apply', stdout=output)
        old.refresh_from_db()
        self.assertEqual(old.birth_date, '2015-06-10', output.getvalue())
        self.assertEqual(old.school, 'Trường A')
        self.assertEqual(old.parent, 'Mai Thu Trang')
        self.assertEqual(set(old.participations.values_list('session_id', flat=True)), {'TESTA-2026', 'TESTB-2026'})
        # The individual registration that already existed keeps its own billing.
        self.assertIsNone(old.participations.get(session_id='TESTA-2026').school_registration_id)
        new = Candidate.objects.get(name='Trần Minh Bình')
        self.assertEqual(new.identity, '001315028059')
        self.assertEqual(new.email, 'b@example.test')

        # The school later corrects a pupil: only --update-profiles overwrites.
        changed = self.appendix_workbook([
            [1, 'Nguyễn Bảo Hoàng', datetime(2015, 6, 10), '001215047377', '7A3', 'TESTA & TESTB', 'Mai Thu Trang', '0984127270', 'new@example.test', '500.000VNĐ'],
        ])
        with open(handle.name, 'wb') as stream:
            stream.write(changed)
        with patch('examination.partner_contact_sync.launch_partner_contact_sync'):
            call_command(*arguments, '--apply', stdout=io.StringIO())
            old.refresh_from_db()
            self.assertEqual(old.class_name, '6A1')
            call_command(*arguments, '--update-profiles', '--apply', stdout=io.StringIO())
        old.refresh_from_db()
        self.assertEqual((old.class_name, old.email), ('7A3', 'new@example.test'))
        self.assertEqual(new.participations.get().school_registration.school, 'Trường A')

    def test_three_contests_split_fee_by_prices_learnt_from_the_file(self):
        ExamSession.objects.create(id='TESTC-2026', competition_id='TESTC', code='TESTC', name='TESTC', parent='TESTC', organizer='Test', time='2026–2027', sort_key='TESTC',
            rounds=[{'id': 'round-1', 'name': 'Vòng 1', 'date': '2026-10-25'}])
        rows = [
            [1, 'Nguyễn Minh An', '12/07/2015', '001215012345', 'Lớp 6', 'TESTA & TESTB & TESTC', '900.000VNĐ', 'Nguyễn Văn A', '0901234567', 'a@example.test', ''],
            [2, 'Trần Minh Bình', '10/02/2014', '002214012345', 'Lớp 7', 'TESTC', '400.000', 'Trần Văn A', '0907654321', 'b@example.test', ''],
            [3, 'Lê Thu Hà', '01/03/2014', '003214012345', 'Lớp 7', 'TESTA', '250000', 'Lê Văn A', '0907654000', 'c@example.test', ''],
        ]
        preview = self.preview(self.workbook(rows))
        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        self.assertEqual(preview.data['summary']['registrations'], 5)
        amounts = {group['sessionId']: group['amount'] for group in preview.data['groups']}
        # TESTA 250k (learnt) + TESTC 400k (learnt) → TESTB gets the remaining 250k.
        self.assertEqual(amounts, {'TESTA-2026': 500000, 'TESTB-2026': 250000, 'TESTC-2026': 800000})

    def test_many_contests_and_separators_without_price_hints(self):
        from .school_import import registered_contests
        for value in ('TESTA & TESTB & TESTC', 'TESTA, TESTB và TESTC', 'TESTA - TESTB - TESTC', 'TESTA | TESTB | TESTC', 'TESTA+TESTB/TESTC'):
            self.assertEqual(registered_contests(value, []), ['TESTA', 'TESTB', 'TESTC'], value)
        from .school_import import split_fee
        self.assertEqual(split_fee(750000, ['TESTA', 'TESTB', 'TESTC'], {}), {'TESTA': 250000, 'TESTB': 250000, 'TESTC': 250000})
        self.assertIsNone(split_fee(800000, ['TESTA', 'TESTB', 'TESTC'], {}))
        self.assertIsNone(split_fee(800000, ['TESTA', 'TESTB'], {'testa': 250000, 'testb': 250000}))

    def test_paid_checkbox_column_is_not_read_as_the_fee(self):
        from .sync import resolve_column_indices
        mapping = resolve_column_indices(['STT', 'Họ và tên thí sinh', 'Nộp lệ phí', 'Ghi chú'], include_defaults=False)
        self.assertNotIn('amount', mapping)
        self.assertEqual(mapping['paymentStatus'], 2)
        mapping = resolve_column_indices(['Họ và tên thí sinh', 'Lệ phí', 'Nộp lệ phí'], include_defaults=False)
        self.assertEqual((mapping['amount'], mapping['paymentStatus']), (1, 2))

    def test_official_appendix_template_reads_school_block_and_text_fees(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Cả trường'
        sheet.append(['DANH SÁCH ĐĂNG KÝ THAM DỰ CUỘC THI'])
        sheet.append([])
        sheet.append(['Nội dung', None, 'Thông tin'])
        sheet.append(['Tên trường học/tổ chức', None, 'Trường A'])
        sheet.append(['Loại hình đơn vị', None, 'THCS'])
        sheet.append(['Địa chỉ', None, 'Phố A, phường B, Hà Nội'])
        sheet.append(['Người phụ trách', None, 'Họ và tên: Người Liên Lạc. Chức vụ: Phó Hiệu trưởng.'])
        sheet.append(['Thông tin liên hệ', None, 'Số điện thoại: 0912345678. ', None, None, 'Email: school@example.test.'])
        sheet.append(['Địa chỉ nhận chứng nhận (nếu có)', None, 'Trường A'])
        sheet.append([])
        sheet.append(['STT', 'Họ và tên thí sinh', 'Ngày, tháng,\nnăm sinh', 'Căn cước công dân/\nHộ chiếu', 'Lớp đang\nhọc', 'Cuộc thi đăng ký\n(FIMO/FIEO)', 'Họ và tên phụ huynh/\nngười giám hộ', 'Số điện thoại', 'Email học sinh/\nphụ huynh', 'Lệ phí', 'Nộp lệ phí', 'Ghi chú', 'Tổng thí sinh'])
        sheet.append([1, 'Nguyễn Minh An', datetime(2015, 6, 10), '001215012345', '6A1', 'TESTA & TESTB', 'Nguyễn Văn A', '0901234567', 'parent@example.test', '500.000VNĐ', True, None, 'TESTA'])
        sheet.append([2, 'Trần Minh Bình', datetime(2015, 2, 14), '1315028059', '6A1', 'TESTB', 'Trần Văn A', '907654321', 'parent2@example.test', '250,000VNĐ', False, None, 'TESTB'])
        stream = io.BytesIO()
        workbook.save(stream)
        preview = self.preview(stream.getvalue(), {'academicYear': '2026-2027'})
        partner = preview.data['partner']
        self.assertEqual(partner['school'], 'Trường A')
        self.assertEqual(partner['representative'], 'Người Liên Lạc')
        self.assertEqual(partner['phone'], '0912345678')
        self.assertEqual(partner['email'], 'school@example.test')
        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        self.assertEqual(preview.data['summary']['candidates'], 2)
        self.assertEqual(preview.data['summary']['registrations'], 3)
        self.assertEqual({group['amount'] for group in preview.data['groups']}, {250000, 500000})
        self.commit(stream.getvalue(), {'academicYear': '2026-2027'})
        an = Candidate.objects.get(name='Nguyễn Minh An')
        self.assertEqual(an.birth_date, '2015-06-10')
        binh = Candidate.objects.get(name='Trần Minh Bình')
        self.assertEqual(binh.identity, '001315028059')
        self.assertEqual(binh.phone, '0907654321')

    def test_missing_fields_warn_without_dropping_students(self):
        rows = [r[:] for r in (self.rows[0], self.rows[2])]
        rows[0][2] = ''
        rows[-1][4] = ''
        preview = self.preview(self.workbook(rows, school_headers=True))

        warnings = [i for i in preview.data['issues'] if i['level'] == 'warning' and 'lệ phí' in i['message']]
        self.assertEqual(len(warnings), 1)
        self.assertIn('2 dòng', warnings[0]['message'])
        self.assertEqual({i['row'] for i in preview.data['issues'] if i['level'] == 'warning' and i['row']}, {2, 3})
        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        result = self.commit(self.workbook(rows, school_headers=True), {'partner': {'school': 'Trường A'}, 'academicYear': '2026-2027'})
        self.assertEqual(result.data['summary']['candidates'], 2)
        self.assertEqual(result.data['summary']['registrations'], 2)
        self.assertEqual(Candidate.objects.get(name='Nguyễn Minh An').birth_date, '')
        self.assertEqual(Candidate.objects.get(name='Trần Minh Bình').class_name, '')
        partner = result.data['partners'][0]
        self.assertFalse(partner.get('phone'))
        self.assertFalse(partner.get('email'))

    def test_leading_zeros_restore_identity_and_phone_and_reuse_existing_candidate(self):
        candidate = Candidate.objects.create(id='EXISTING', code='EXISTING', name='Nguyễn Minh An', birth_date='2015-07-12', identity='1215012345', phone='901234567', school='Trường A', sort_key='a')
        row = self.rows[0][:]
        row[2] = ''
        row[3] = 1215012345
        row[8] = 901234567
        preview = self.preview(self.workbook([row]))
        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        self.assertEqual(preview.data['summary']['newCandidates'], 0)
        self.assertEqual(preview.data['profiles'][0]['profile']['identity'], '001215012345')
        self.assertEqual(preview.data['profiles'][0]['profile']['phone'], '0901234567')
        self.commit(self.workbook([row]))
        self.assertEqual(Candidate.objects.count(), 1)
        candidate.refresh_from_db()
        self.assertEqual(candidate.identity, '001215012345')
        self.assertEqual(candidate.phone, '0901234567')
        self.assertEqual(candidate.birth_date, '2015-07-12')

    def test_identifier_formatting_preserves_blanks_passports_and_long_values(self):
        from .sync import format_identity, format_phone
        for raw, expected in [('', ''), ('B1234567', 'B1234567'), ('123456789012', '123456789012'), ('1234567890123', '1234567890123')]:
            self.assertEqual(format_identity(raw), expected)
        for raw, expected in [('', ''), ('901234567', '0901234567'), ('+84 901 234 567', '0901234567'), ('0084901234567', '0901234567'), ('0901234567', '0901234567'), ('0901234567,', '0901234567'), ('12345678901', '12345678901'), ('0901234567 / 0907654321', '0901234567 / 0907654321')]:
            self.assertEqual(format_phone(raw), expected)

    def test_repeated_student_with_one_blank_birth_date_reuses_complete_row(self):
        rows = [r[:] for r in self.rows[:2]]
        rows[0][2] = ''
        result = self.commit(self.workbook(rows))
        self.assertEqual(result.data['summary']['candidates'], 1)
        self.assertEqual(result.data['summary']['registrations'], 2)
        self.assertEqual(Candidate.objects.get().birth_date, '2015-07-12')

    def test_unknown_combined_contest_skips_the_whole_row(self):
        row = self.rows[0][:]
        row[5] = 'TESTA & UNKNOWN'
        preview = self.preview(self.workbook([row], school_headers=True))

        # Half of a row is never imported; with no other row nothing commits.
        self.assertFalse(preview.data['canCommit'])
        self.assertEqual(preview.data['skippedRows'], [2])
        self.assertEqual(CandidateParticipation.objects.count(), 0)

    def test_exact_session_name_with_separator_is_preserved(self):
        session = self.sessions['TESTA']
        session.name = 'Toán / Khoa học'
        session.save()
        row = self.rows[0][:]
        row[5] = session.name
        preview = self.preview(self.workbook([row]))

        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        self.assertEqual(preview.data['summary']['registrations'], 1)
        self.assertEqual(preview.data['routes'][0]['sessionId'], session.pk)

    def test_retry_and_exact_duplicate_do_not_double_bill(self):
        rows = self.rows + [self.rows[0][:]]
        rows[-1][0] = 4
        content = self.workbook(rows)
        self.commit(content)
        billing = ExaminationBillingRecord.objects.first()
        billing.transfer_status = 'confirmed'
        billing.invoice_number = 'HD-001'
        billing.save()
        self.commit(content)
        self.assertEqual(Candidate.objects.count(), 2)
        self.assertEqual(CandidateParticipation.objects.count(), 3)
        self.assertEqual(ExaminationBillingRecord.objects.count(), 2)
        billing.refresh_from_db()
        self.assertEqual(billing.transfer_status, 'confirmed')
        self.assertEqual(billing.invoice_number, 'HD-001')
        again = self.commit(content)
        self.assertEqual(again.data['summary']['newCandidates'], 0)
        self.assertEqual(again.data['summary']['existingCandidates'], 2)
        self.assertEqual(again.data['summary']['newRegistrations'], 0)
        self.assertEqual(again.data['summary']['existingRegistrations'], 3)

    def test_metadata_and_duplicate_tabs_use_one_selected_table(self):
        content = self.workbook(metadata=True, copy=True)
        raw, metadata, sheet, sheets = read_workbook(content)
        self.assertEqual(metadata['school'], 'Trường A')
        self.assertEqual(metadata['email'], 'school@example.test')
        self.assertEqual(len(raw), 3)
        self.assertEqual(sheet, 'Đăng ký tham dự')
        self.assertEqual(len(sheets), 2)
        self.commit(content, {'academicYear': '2026-2027'})

    def individual_registration(self, paid=True):
        candidate = Candidate.objects.create(id='EXISTING', code='EXISTING', name='Nguyễn Minh An', birth_date='2015-07-12', identity='001215012345', school='Trường A', class_name='', sort_key='a')
        participation = CandidateParticipation.objects.create(candidate=candidate, session=self.sessions['TESTA'], registration_method='Cá nhân', source='Google Form', registration_data={'paymentProof': 'original-proof', 'generalNote': 'Original'})
        billing = participation.billing
        if paid:
            billing.amount = 250000
            billing.transfer_status = 'confirmed'
            billing.transfer_reference = 'original-transfer'
            billing.invoice_status = 'issued'
            billing.invoice_number = 'original-invoice'
            billing.seen_by_accountant = True
            billing.save()
        return participation

    def test_paid_individual_is_preserved_and_only_missing_contests_are_added(self):
        participation = self.individual_registration()
        room = ExamRoom.objects.create(session=self.sessions['TESTA'], round_id='round-1', round_name='Vòng 1', label='Original room', room_number='1', mode='IN_PERSON', capacity=10)
        round_result = RoundResult.objects.create(participation=participation, round_id='round-1', round_name='Vòng 1', exam_room=room, room_name='Original room', exam_date='2026-10-25', eligibility='Đủ điều kiện')
        proof = TransferProof.objects.create(billing=participation.billing, session=self.sessions['TESTA'], image=b'original', image_type='image/png', filename='original.png', created_by=self.finance.email)
        models = [(CandidateParticipation, participation.pk), (ExaminationBillingRecord, participation.billing.pk), (RoundResult, round_result.pk), (TransferProof, proof.pk)]
        snapshots = [model.objects.filter(pk=pk).values().get() for model, pk in models]
        preview = self.preview()
        self.assertTrue(preview.data['canCommit'], preview.data['issues'])
        self.assertEqual(preview.data['summary']['newCandidates'], 1)
        self.assertEqual(preview.data['summary']['preservedIndividualRegistrations'], 1)
        self.assertEqual(preview.data['summary']['newRegistrations'], 2)
        report = {g['competitionCode']: g for g in preview.data['groups']}
        self.assertEqual(report['TESTA']['amount'], 250000)
        self.assertEqual(report['TESTA']['alreadyAssigned'], 1)
        self.commit()
        self.assertEqual(Candidate.objects.count(), 2)
        self.assertEqual(CandidateParticipation.objects.count(), 3)
        for (model, pk), snapshot in zip(models, snapshots):
            self.assertEqual(model.objects.filter(pk=pk).values().get(), snapshot)
        participation.candidate.refresh_from_db()
        self.assertEqual(participation.candidate.class_name, 'Lớp 6')
        self.assertEqual(ExaminationBillingRecord.objects.get(school_registration__session=self.sessions['TESTA']).amount, 250000)
        self.assertEqual(ExaminationBillingRecord.objects.get(school_registration__session=self.sessions['TESTB']).amount, 450000)
        again = self.commit()
        self.assertEqual(again.data['summary']['newRegistrations'], 0)
        self.assertEqual(again.data['summary']['preservedIndividualRegistrations'], 1)
        self.assertEqual(CandidateParticipation.objects.count(), 3)
        self.assertEqual(ExaminationBillingRecord.objects.count(), 3)

    def test_only_existing_individual_creates_no_school_bill_or_room_assignment(self):
        for paid in (False, True):
            with self.subTest(paid=paid):
                if not Candidate.objects.exists():
                    participation = self.individual_registration(paid=paid)
                else:
                    participation.billing.amount = 250000
                    participation.billing.transfer_status = 'confirmed'
                    participation.billing.save()
                room = ExamRoom.objects.create(session=self.sessions['TESTA'], round_id='round-1', round_name='Vòng 1', label='Available room', room_number=str(paid), mode='IN_PERSON', capacity=10)
                snapshot = CandidateParticipation.objects.filter(pk=participation.pk).values().get()
                result = self.commit(self.workbook([self.rows[0]]))
                self.assertEqual(result.data['summary']['newRegistrations'], 0)
                self.assertEqual(result.data['summary']['preservedIndividualRegistrations'], 1)
                self.assertEqual(result.data['report'][0]['assigned'], 0)
                self.assertEqual(result.data['report'][0]['waiting'], 0)
                self.assertEqual(len(result.data['sessions']), 1)
                self.assertEqual(SchoolRegistration.objects.count(), 0)
                self.assertEqual(ExaminationBillingRecord.objects.count(), 1)
                self.assertEqual(room.assignments.count(), 0)
                self.assertEqual(CandidateParticipation.objects.filter(pk=participation.pk).values().get(), snapshot)

    def test_individual_changed_since_preview_requires_new_preview(self):
        participation = self.individual_registration()
        content = self.workbook()
        preview = self.preview(content)
        participation.general_note = 'Changed after preview'
        participation.save()
        response = self.post(content, self.options | {'action': 'commit', 'previewToken': preview.data['previewToken']})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(Candidate.objects.count(), 1)
        self.assertEqual(CandidateParticipation.objects.count(), 1)

    def test_faulty_rows_are_skipped_and_other_pupils_import(self):
        rows = [r[:] for r in self.rows]
        rows[0][5] = 'UNKNOWN'
        rows[1][2] = '31/02/2015'
        content = self.workbook(rows)
        preview = self.preview(content)
        self.assertTrue(preview.data['canCommit'])
        self.assertEqual(preview.data['skippedRows'], [2, 3])
        self.assertTrue(any(issue['row'] == 2 and issue['level'] == 'error' for issue in preview.data['issues']))
        self.commit(content)
        self.assertEqual(list(Candidate.objects.values_list('name', flat=True)), ['Trần Minh Bình'])

    def test_ambiguous_sessions_require_mapping(self):
        original = self.sessions['TESTA']
        ExamSession.objects.create(id='TESTA-other', competition_id='TESTA', code='TESTA', name='Đợt khác', parent='TESTA', organizer='Test', time=original.time, sort_key='z', rounds=original.rounds)
        preview = self.preview()
        self.assertEqual(preview.data['skippedRows'], [2, 4])
        self.commit(options=self.options | {'sessionMapping': {'TESTA': original.pk}})
        self.assertEqual(CandidateParticipation.objects.filter(session_id=original.pk).count(), 2)

    def test_stale_preview_and_tampered_file_are_rejected(self):
        content = self.workbook()
        preview = self.preview(content)
        self.sessions['TESTA'].note = 'Changed after preview'
        self.sessions['TESTA'].save()
        response = self.post(content, self.options | {'action': 'commit', 'previewToken': preview.data['previewToken']})
        self.assertEqual(response.status_code, 409)
        preview = self.preview(content)
        rows = [r[:] for r in self.rows]
        rows[0][6] = 999000
        response = self.post(self.workbook(rows), self.options | {'action': 'commit', 'previewToken': preview.data['previewToken']})
        self.assertEqual(response.status_code, 409)
        self.assertEqual(Candidate.objects.count(), 0)

    def test_existing_school_contact_auto_fills_and_profile_is_reused(self):
        self.commit()
        partners = SystemConfig.objects.get(key='examination_partners').data['partners']
        self.assertEqual(len(partners), 1)
        response = self.preview(options={'partnerId': partners[0]['id'], 'academicYear': '2026-2027'})
        self.assertEqual(response.data['partner']['email'], 'school@example.test')
        self.assertEqual(response.data['summary']['newCandidates'], 0)

    def test_conflicting_duplicate_fees_block_commit(self):
        rows = self.rows + [self.rows[0][:]]
        rows[-1][6] = 999000
        preview = self.preview(self.workbook(rows))
        self.assertIn(5, preview.data['skippedRows'])

    def test_same_identity_with_different_names_in_file_is_not_duplicated(self):
        rows = [r[:] for r in self.rows]
        rows[1][1] = 'Nguyễn Minh Anh'
        preview = self.preview(self.workbook(rows))
        self.assertIn(3, preview.data['skippedRows'])
        self.assertTrue(any('Cùng giấy tờ' in issue['message'] for issue in preview.data['issues']))

    def test_reverse_sheet_sync_preserves_school_accounting_metadata(self):
        from .views import upsert_participation_history
        self.commit()
        participation = CandidateParticipation.objects.first()
        original = dict(participation.registration_data)
        upsert_participation_history(participation.candidate, participation.session_id, [], 'Sheet', {'generalNote': 'Updated from Sheet'})
        participation.refresh_from_db()
        self.assertEqual(participation.registration_data['schoolFee'], original['schoolFee'])
        self.assertEqual(participation.registration_data['schoolPartnerId'], original['schoolPartnerId'])
        self.assertTrue(self.preview().data['canCommit'])

    def test_shared_parent_contact_with_different_identity_requires_confirmation(self):
        Candidate.objects.create(id='EXISTING', code='EXISTING', name='Nguyễn Minh An', birth_date='2015-07-12', identity='009999012345', phone='0901234567', email='parent@example.test', school='Trường A', class_name='Lớp 6', sort_key='a')
        preview = self.preview()
        self.assertEqual(preview.data['skippedRows'], [2, 3])
        options = self.options | {'candidateMatches': {'2': '__new__', '3': '__new__'}}
        self.commit(options=options)
        self.assertEqual(Candidate.objects.filter(name='Nguyễn Minh An').count(), 2)

    def test_settled_school_account_cannot_silently_add_new_fees(self):
        self.commit()
        billing = ExaminationBillingRecord.objects.get(school_registration__session=self.sessions['TESTA'])
        billing.transfer_status = 'confirmed'
        billing.save()
        rows = self.rows + [[4, 'Lê Minh Chi', '10/02/2014', '003214012345', 'Lớp 7', 'TESTA', 250000, 'Lê Văn A', '0907654322', 'parent3@example.test', '']]
        self.assertFalse(self.preview(self.workbook(rows)).data['canCommit'])

    def test_allocation_respects_existing_assignments_and_capacity(self):
        room = ExamRoom.objects.create(session=self.sessions['TESTA'], round_id='round-1', round_name='Vòng 1', common_name='Phòng', room_number='1', label='Phòng 1', mode='IN_PERSON', location='Trường A', capacity=1, allocation_strategy='CAPACITY')
        preview = self.preview()
        group = next(g for g in preview.data['groups'] if g['sessionId'] == self.sessions['TESTA'].pk)
        self.assertEqual(group['assigned'], 1)
        self.assertEqual(group['waiting'], 1)
        self.commit()
        self.assertEqual(RoundResult.objects.filter(exam_room=room).count(), 1)
        self.commit()
        self.assertEqual(RoundResult.objects.filter(exam_room=room).count(), 1)

    def test_transaction_rolls_back_candidate_partner_and_billing(self):
        content = self.workbook()
        preview = self.preview(content)
        with patch('examination.views.upsert_participation_history', side_effect=ValueError('Injected failure')):
            response = self.post(content, self.options | {'action': 'commit', 'previewToken': preview.data['previewToken']})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(Candidate.objects.count(), 0)
        self.assertEqual(SchoolRegistration.objects.count(), 0)
        self.assertEqual(ExaminationBillingRecord.objects.count(), 0)
        self.assertEqual(CandidateSheetOutbox.objects.count(), 0)

    def test_finance_uploads_multiple_images_append_and_access_is_private(self):
        self.commit()
        billing = ExaminationBillingRecord.objects.first()
        png = lambda name: SimpleUploadedFile(name, b'\x89PNG\r\n\x1a\n' + b'x' * 20, content_type='image/png')
        endpoint = f'/api/examination/billing/records/{billing.pk}/proofs'
        self.assertEqual(self.client.post(endpoint, {'images': [png('1.png')]}, format='multipart').status_code, 403)
        self.client.force_authenticate(self.finance)
        response = self.client.post(endpoint, {'images': [png('1.png'), png('2.png')]}, format='multipart')
        self.assertEqual(response.status_code, 200, response.data)
        self.assertEqual(len(response.data['proofs']), 2)
        self.client.post(endpoint, {'images': [png('3.png')]}, format='multipart')
        self.assertEqual(billing.proofs.count(), 3)
        self.assertEqual(set(billing.proofs.values_list('session_id', flat=True)), {billing.school_registration.session_id})
        invalid = self.client.post(endpoint, {'images': [png('4.png'), SimpleUploadedFile('invalid.png', b'not an image')]}, format='multipart')
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(billing.proofs.count(), 3)
        proof = billing.proofs.first()
        self.client.force_authenticate(None)
        self.assertNotEqual(self.client.get(f'/api/examination/billing/proofs/{proof.pk}/image').status_code, 200)

    def test_unmatched_images_wait_for_verified_contest(self):
        self.commit()
        self.client.force_authenticate(self.finance)
        png = lambda name: SimpleUploadedFile(name, b'\x89PNG\r\n\x1a\n' + b'x' * 20, content_type='image/png')
        response = self.client.post('/api/examination/billing/unmatched', {'amount': '250000', 'images': [png('1.png'), png('2.png')]}, format='multipart')
        self.assertEqual(response.status_code, 201)
        self.assertEqual(len(response.data['proofs']), 2)
        self.assertTrue(all(p['sessionId'] is None for p in response.data['proofs']))
        participation = CandidateParticipation.objects.first()
        self.client.force_authenticate(self.user)
        resolved = self.client.post(f'/api/examination/billing/unmatched/{response.data["id"]}/resolve', {'candidateCode': participation.candidate.code, 'competitionCode': participation.session.code, 'resolutionNote': 'Đã xác minh'}, format='json')
        self.assertEqual(resolved.status_code, 200)
        self.assertEqual({p['sessionId'] for p in resolved.data['proofs']}, {participation.session_id})

    def test_drive_is_disabled_until_configured_then_routes_and_retries(self):
        self.commit()
        billing = ExaminationBillingRecord.objects.first()
        proof = TransferProof.objects.create(billing=billing, session=billing.school_registration.session, image=b'png', image_type='image/png', filename='proof.png', created_by='test')
        service = MagicMock()
        self.assertFalse(archive_pending(service=service)['enabled'])
        service.files.assert_not_called()
        config = SystemConfig.objects.create(key='examination_proof_drive', data={'enabled': True, 'competitionFolders': {proof.session.code: 'competition-folder'}, 'sessionFolders': {proof.session_id: 'session-folder'}})
        self.assertEqual(folder_for(proof.session), 'session-folder')
        service.files.return_value.list.return_value.execute.return_value = {'files': [{'id': 'existing-file'}]}
        self.assertEqual(archive_pending(service=service)['synced'], 1)
        proof.refresh_from_db()
        self.assertEqual(proof.drive_folder_id, 'session-folder')
        self.assertEqual(proof.drive_file_id, 'existing-file')
        service.files.return_value.create.assert_not_called()
        self.assertEqual(archive_pending(service=service)['synced'], 0)
