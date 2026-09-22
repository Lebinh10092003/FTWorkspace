"""Word export for the weekly-report preview.

The report follows the supplied Fermat weekly report sample: an A4 page with a
centred three-line title followed by the three narrative sections.  It is built
from structured data rather than copying any employee's historical report.
"""

from io import BytesIO
import re
import unicodedata

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


FONT = "Times New Roman"


def _safe_text(value):
    return str(value or "").strip()


def _items(value):
    if not isinstance(value, list):
        return []
    return [_safe_text(item) for item in value if _safe_text(item)]


def _style(run, *, bold=False, size=13):
    run.bold = bold
    run.font.name = FONT
    run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _paragraph(document, text="", *, bold=False, size=13, align=None, before=0, after=2):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    if align is not None:
        paragraph.alignment = align
    _style(paragraph.add_run(text), bold=bold, size=size)
    return paragraph


def _section(document, heading, items):
    _paragraph(document, heading, bold=True, before=5, after=2)
    rows = _items(items) or ["Không có nội dung được ghi nhận."]
    for item in rows:
        _paragraph(document, f"- {item}", after=1)


def _filename(report):
    raw = _safe_text(report.employee_name) or "nhan-vien"
    normalized = "".join(
        char for char in unicodedata.normalize("NFD", raw)
        if unicodedata.category(char) != "Mn"
    ).replace("Đ", "D").replace("đ", "d")
    slug = re.sub(r"[^A-Za-z0-9]+", "-", normalized).strip("-").lower() or "nhan-vien"
    return f"Bao-cao-tuan-{report.completed_week}-ke-hoach-tuan-{report.planned_week}-{slug}.docx"


def build_weekly_report_docx(report):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = Cm(2)
    section.right_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(13)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    title = _paragraph(document, "", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER, after=7)
    for index, line in enumerate((
        "BÁO CÁO",
        f"Kết quả tuần {report.completed_week}, nhiệm vụ dự kiến tuần {report.planned_week}",
        f"({report.employee_name})",
    )):
        run = title.add_run(line)
        _style(run, bold=True, size=14)
        if index < 2:
            run.add_break()

    _section(document, f"1. Công việc đã thực hiện tuần {report.completed_week}", report.completed_items)
    _section(document, "2. Tồn tại, khó khăn, vướng mắc", report.difficulties)
    _section(document, f"3. Nhiệm vụ tuần {report.planned_week}", report.planned_items)

    stream = BytesIO()
    document.save(stream)
    return _filename(report), stream.getvalue()
