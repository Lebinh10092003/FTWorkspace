"""Word (.docx) export builder for news posts following FermatTech template."""

from __future__ import annotations

import io
import re
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from documents.filename_utils import sanitize_docx_filename
from .models import NewsPost
from .validator import format_hashtags_string, split_contact_info_lines

FONT_PRIMARY = "Aptos"
FONT_FALLBACK = "Calibri"


def _set_run_font(
    run,
    *,
    font_name: str = FONT_PRIMARY,
    size_pt: float = 11,
    bold: bool = False,
    italic: bool = False,
) -> None:
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

    # Đảm bảo font hiển thị tiếng Việt và ký tự đặc biệt/emoji đúng
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), font_name)
    rFonts.set(qn("w:hAnsi"), font_name)
    rFonts.set(qn("w:eastAsia"), font_name)
    rFonts.set(qn("w:cs"), font_name)


def _add_paragraph(
    document: Document,
    text: str = "",
    *,
    bold: bool = False,
    italic: bool = False,
    size_pt: float = 11,
    align: WD_ALIGN_PARAGRAPH = WD_ALIGN_PARAGRAPH.LEFT,
    space_before: float = 0,
    space_after: float = 4,
    line_spacing: float = 1.15,
):
    p = document.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        run = p.add_run(text)
        _set_run_font(run, size_pt=size_pt, bold=bold, italic=italic)
    return p


def _parse_date(date_str: str) -> Optional[date]:
    """Parse chuỗi ngày tháng dạng YYYY-MM-DD, YYYY.MM.DD, DD/MM/YYYY, etc."""
    if not date_str or not isinstance(date_str, str):
        return None
    s = date_str.strip()
    # YYYY-MM-DD, YYYY.MM.DD, YYYY/MM/DD
    m_iso = re.search(r"\b(20\d{2})[./-](\d{1,2})[./-](\d{1,2})\b", s)
    if m_iso:
        y, m, d = int(m_iso.group(1)), int(m_iso.group(2)), int(m_iso.group(3))
        try:
            return date(y, m, d)
        except ValueError:
            pass
    # DD/MM/YYYY, DD-MM-YYYY, DD.MM.YYYY
    m_vn = re.search(r"\b(\d{1,2})[./-](\d{1,2})[./-](20\d{2})\b", s)
    if m_vn:
        d, m, y = int(m_vn.group(1)), int(m_vn.group(2)), int(m_vn.group(3))
        try:
            return date(y, m, d)
        except ValueError:
            pass
    # ngày D[D] tháng M[M] năm YYYY
    m_text = re.search(r"ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(20\d{2})", s, flags=re.IGNORECASE)
    if m_text:
        d, m, y = int(m_text.group(1)), int(m_text.group(2)), int(m_text.group(3))
        try:
            return date(y, m, d)
        except ValueError:
            pass
    return None


def _find_dates_in_text(text: str) -> list[date]:
    """Tìm tất cả ngày tháng trong văn bản."""
    if not text or not isinstance(text, str):
        return []
    dates: list[date] = []
    # 1. YYYY-MM-DD / YYYY.MM.DD / YYYY/MM/DD
    for m in re.finditer(r"\b(20\d{2})[./-](\d{1,2})[./-](\d{1,2})\b", text):
        try:
            dates.append(date(int(m.group(1)), int(m.group(2)), int(m.group(3))))
        except ValueError:
            pass
    # 2. DD/MM/YYYY / DD-MM-YYYY / DD.MM.YYYY
    for m in re.finditer(r"\b(\d{1,2})[./-](\d{1,2})[./-](20\d{2})\b", text):
        try:
            dates.append(date(int(m.group(3)), int(m.group(2)), int(m.group(1))))
        except ValueError:
            pass
    # 3. ngày D[D] tháng M[M] năm YYYY
    for m in re.finditer(r"ngày\s+(\d{1,2})\s+tháng\s+(\d{1,2})\s+năm\s+(20\d{2})", text, flags=re.IGNORECASE):
        try:
            dates.append(date(int(m.group(3)), int(m.group(2)), int(m.group(1))))
        except ValueError:
            pass
    return dates


def _extract_event_date(news_post: NewsPost) -> date:
    """Lấy ngày diễn ra sự kiện/buổi làm việc (lấy từ dữ liệu đầu vào/tiến trình nếu có,
    nhiều buổi thì ngày buổi cuối); không có thì ngày tạo bài (ICT Asia/Ho_Chi_Minh)."""
    input_payload = news_post.input_payload if isinstance(news_post.input_payload, dict) else {}
    dates: list[date] = []

    # 1. Từ danh sách sessions (lịch các buổi)
    sessions = input_payload.get("sessions")
    if isinstance(sessions, list):
        for s in sessions:
            if isinstance(s, dict):
                d_str = s.get("session_date") or s.get("sessionDate") or s.get("date")
                parsed = _parse_date(str(d_str or ""))
                if parsed:
                    dates.append(parsed)

    # 2. Tìm trong agenda_files (tiến trình)
    if not dates:
        agenda_files = input_payload.get("agenda_files") or input_payload.get("agendaFiles") or []
        if isinstance(agenda_files, list):
            for f in agenda_files:
                if isinstance(f, dict):
                    dates.extend(_find_dates_in_text(str(f.get("text") or "")))
                    dates.extend(_find_dates_in_text(str(f.get("name") or "")))

    # 3. Tìm trong notes, event_date hoặc title
    if not dates:
        notes = str(input_payload.get("notes") or "")
        dates.extend(_find_dates_in_text(notes))
        evt_date = input_payload.get("event_date") or input_payload.get("eventDate")
        parsed = _parse_date(str(evt_date or ""))
        if parsed:
            dates.append(parsed)
        dates.extend(_find_dates_in_text(str(news_post.title or "")))

    # Nhiều buổi thì ngày buổi cuối (ngày lớn nhất)
    if dates:
        return max(dates)

    # Không có thì ngày tạo bài (ICT Asia/Ho_Chi_Minh)
    if news_post.created_at:
        try:
            local_dt = timezone.localtime(news_post.created_at)
        except Exception:
            local_dt = news_post.created_at
        return local_dt.date()

    return timezone.localtime(timezone.now()).date()


def _shorten_unit_name(unit: str) -> str:
    """Rút gọn tên cấp học và đơn vị theo chuẩn công ty:
    Trường Tiểu học & THCS Lý Thường Kiệt -> TH, THCS Lý Thường Kiệt
    Trường THCS Giảng Võ -> THCS Giảng Võ
    Trường Mầm non Vĩnh Hưng -> MN Vĩnh Hưng
    """
    if not unit:
        return ""
    text = str(unit).strip()
    # Bỏ emoji và ký tự điều khiển
    text = re.sub(r"[\U00010000-\U0010ffff\u2600-\u27bf\u2300-\u23ff\u2b50-\u2b55\ufe00-\ufe0f\u200d\u20e3]", " ", text)
    # Bỏ dấu hai chấm
    text = text.replace(":", " ")
    # Bỏ chữ Trường / trường liên cấp
    text = re.sub(r"\b(?:trường\s+liên\s+cấp|trường)\s+", "", text, flags=re.IGNORECASE)
    # Rút gọn các cấp học ghép
    text = re.sub(r"\b(?:tiểu\s*học|th)\s*(?:&|và|\+)\s*(?:trung\s*học\s*cơ\s*sở|thcs)\b", "TH, THCS", text, flags=re.IGNORECASE)
    text = re.sub(r"\b(?:trung\s*học\s*cơ\s*sở|thcs)\s*(?:&|và|\+)\s*(?:trung\s*học\s*phổ\s*thông|thpt)\b", "THCS, THPT", text, flags=re.IGNORECASE)
    # Rút gọn từng cấp học đơn lẻ
    text = re.sub(r"\btrung\s*học\s*cơ\s*sở\b", "THCS", text, flags=re.IGNORECASE)
    text = re.sub(r"\btrung\s*học\s*phổ\s*thông\b", "THPT", text, flags=re.IGNORECASE)
    text = re.sub(r"\btiểu\s*học\b", "TH", text, flags=re.IGNORECASE)
    text = re.sub(r"\bmầm\s*non\b", "MN", text, flags=re.IGNORECASE)
    text = re.sub(r"\bmẫu\s*giáo\b", "MG", text, flags=re.IGNORECASE)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _extract_newspaper_name(news_post: NewsPost) -> str:
    """Trích xuất tên báo (nếu có) cho bài báo chí (press_article)."""
    input_payload = news_post.input_payload if isinstance(news_post.input_payload, dict) else {}
    direct = str(
        input_payload.get("newspaper_name")
        or input_payload.get("newspaperName")
        or input_payload.get("press_name")
        or input_payload.get("pressName")
        or ""
    ).strip()
    if direct:
        return direct if direct.lower().startswith("báo ") else f"Báo {direct}"

    candidates = [
        str(news_post.title or ""),
        str(input_payload.get("notes") or ""),
    ]
    if isinstance(news_post.content_payload, dict):
        candidates.append(str(news_post.content_payload.get("headline") or ""))

    for c in candidates:
        if not c:
            continue
        m_gd = re.search(r"\b(báo\s+giáo\s+dục\s*(?:và|&)\s*thời\s+đại)\b", c, flags=re.IGNORECASE)
        if m_gd:
            return "Báo Giáo dục và Thời đại"
        m = re.search(r"\b(?:gửi\s+)?(báo\s+[A-ZÀ-Ỹa-zà-ỹ0-9\s&]+?)(?:[,.;\n]|$)", c, flags=re.IGNORECASE)
        if m:
            name = m.group(1).strip()
            name = re.sub(r"\s+(?:ngày|vào|lúc|tại).*$", "", name, flags=re.IGNORECASE).strip()
            if len(name.split()) >= 2:
                return name
    return ""


def _extract_program_or_topic(news_post: NewsPost) -> str:
    """Trích xuất loại hình / chương trình (BNDC, GCE1, GCE, ...)."""
    input_payload = news_post.input_payload if isinstance(news_post.input_payload, dict) else {}
    direct = str(
        input_payload.get("program")
        or input_payload.get("topic")
        or input_payload.get("training_type")
        or ""
    ).strip()
    if direct:
        return direct

    text = " ".join([
        str(news_post.title or ""),
        str(input_payload.get("notes") or ""),
        str(input_payload.get("unit_name") or ""),
    ])

    if re.search(r"\b(?:bndc|bộ\s*nhớ\s*dùng\s*chung)\b", text, flags=re.IGNORECASE):
        return "BNDC"
    if re.search(r"\bgce1\b", text, flags=re.IGNORECASE):
        return "GCE1"
    if re.search(r"\bgce\b", text, flags=re.IGNORECASE):
        return "GCE"
    return ""


def make_safe_docx_filename(news_post: NewsPost) -> str:
    """Tạo tên tệp Word chuẩn theo quy định công ty:
    YYYY.MM.DD. <Loại> <nội dung>.docx

    - Ngày dấu chấm, sau ngày là '. ' (chấm + cách).
    - Giữ tiếng Việt có dấu và khoảng trắng, không slug, không gạch nối thay khoảng trắng.
    - Ngày = ngày diễn ra sự kiện/buổi làm việc (nhiều buổi thì ngày buổi cuối); không có thì ngày tạo bài (ICT).
    - post_type == fb_training: Bài đăng FB tập huấn <nội dung>
    - post_type == fb_meeting: Bài đăng FB <nội dung> (vd: họp/làm việc/cài đặt/giới thiệu sản phẩm...)
    - post_type == press_article: Tin bài <đơn vị/nội dung> gửi <tên báo> nếu có tên báo, không thì Tin bài <nội dung>.
    - <nội dung> = loại hình + cấp học + tên đơn vị rút gọn, ví dụ 'BNDC TH, THCS Lý Thường Kiệt'.
    """
    event_date = _extract_event_date(news_post)
    date_prefix = f"{event_date.year:04d}.{event_date.month:02d}.{event_date.day:02d}. "

    input_payload = news_post.input_payload if isinstance(news_post.input_payload, dict) else {}
    raw_unit = str(
        input_payload.get("unit_name")
        or input_payload.get("unitName")
        or ""
    ).strip()

    title_str = str(news_post.title or "").strip()

    # Rút gọn tên đơn vị
    unit_short = _shorten_unit_name(raw_unit)
    if not unit_short and title_str:
        candidate = re.sub(
            r"^(?:bài\s+đăng\s+fb|tin\s+bài(?:\s+truyền\s+thông)?(?:\s+tại)?|tập\s+huấn(?:\s+bndc)?|buổi\s+làm\s+việc(?:\s+với)?)\s+",
            "",
            title_str,
            flags=re.IGNORECASE,
        ).strip()
        unit_short = _shorten_unit_name(candidate)

    program = _extract_program_or_topic(news_post)
    if program and unit_short and not re.search(rf"\b{re.escape(program)}\b", unit_short, flags=re.IGNORECASE):
        content_unit = f"{program} {unit_short}".strip()
    else:
        content_unit = unit_short or program or ""

    post_type = getattr(news_post, "post_type", NewsPost.TYPE_FB_TRAINING)

    if post_type == NewsPost.TYPE_FB_TRAINING:
        clean_content = re.sub(r"^(?:bài\s+đăng\s+fb\s+)?tập\s+huấn\s*", "", content_unit, flags=re.IGNORECASE).strip()
        body = f"Bài đăng FB tập huấn {clean_content}" if clean_content else "Bài đăng FB tập huấn"

    elif post_type == NewsPost.TYPE_FB_MEETING:
        action_found = ""
        clean_title = re.sub(r"^(?:bài\s+đăng\s+fb|tin\s+bài|buổi)\s+", "", title_str, flags=re.IGNORECASE).strip()
        for kw in ("họp", "làm việc", "cài đặt", "giới thiệu sản phẩm", "giới thiệu"):
            if re.search(rf"\b{re.escape(kw)}\b", clean_title, flags=re.IGNORECASE):
                sub_title = re.sub(r"^(?:buổi\s+)?(?:làm\s+việc|họp|gặp\s+gỡ|trao\s+đổi)\s+(?:với\s+)?", f"{kw} ", clean_title, flags=re.IGNORECASE)
                sub_title = _shorten_unit_name(sub_title)
                if sub_title:
                    action_found = sub_title[0].lower() + sub_title[1:]
                break

        if action_found:
            body = f"Bài đăng FB {action_found}"
        elif content_unit:
            body = f"Bài đăng FB làm việc {content_unit}"
        else:
            body = "Bài đăng FB làm việc"

        body = re.sub(r"^(?:bài\s+đăng\s+fb\s+)+", "Bài đăng FB ", body, flags=re.IGNORECASE).strip()

    elif post_type == NewsPost.TYPE_PRESS_ARTICLE:
        newspaper = _extract_newspaper_name(news_post)
        press_content = unit_short or re.sub(r"^tin\s+bài\s*", "", content_unit, flags=re.IGNORECASE).strip()
        if not press_content:
            press_content = _shorten_unit_name(title_str)
            press_content = re.sub(r"^tin\s+bài\s*", "", press_content, flags=re.IGNORECASE).strip()
        if not press_content:
            press_content = "truyền thông"

        if newspaper:
            body = f"Tin bài {press_content} gửi {newspaper}"
        else:
            body = f"Tin bài {press_content}"
    else:
        body = f"Tin bài {content_unit}" if content_unit else "Tin bài"

    raw_filename = f"{date_prefix}{body}.docx"
    return sanitize_docx_filename(raw_filename, default_stem=f"Tin-bai-{news_post.pk or 'moi'}")


def build_news_post_docx(news_post: NewsPost) -> io.BytesIO:
    """Tạo file Word (.docx) từ nội dung bản tin theo chuẩn bài FB Lý Thường Kiệt / Báo chí."""
    document = Document()

    # Cấu hình khổ giấy A4 và lề trang
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.0)

    # Style mặc định Normal
    normal_style = document.styles["Normal"]
    normal_style.font.name = FONT_PRIMARY
    normal_style.font.size = Pt(11)
    normal_style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_PRIMARY)

    cp: Dict[str, Any] = news_post.content_payload if isinstance(news_post.content_payload, dict) else {}
    post_type = news_post.post_type

    # 1. Tiêu đề in hoa đậm (Headline)
    headline = str(cp.get("headline") or news_post.title or "BÀI ĐĂNG FERMATTECH").strip()
    _add_paragraph(
        document,
        headline.upper(),
        bold=True,
        size_pt=13.5,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=0,
        space_after=8,
    )

    # 2. Đoạn mở đầu (Lead in)
    lead_in = str(cp.get("lead_in") or "").strip()
    if lead_in:
        _add_paragraph(
            document,
            lead_in,
            bold=False,
            size_pt=11,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            space_before=0,
            space_after=6,
        )

    # 3. Các mục nội dung (Sections)
    sections = cp.get("sections") or []
    if isinstance(sections, list):
        for s in sections:
            if not isinstance(s, dict):
                continue
            h = str(s.get("heading") or "").strip()
            b = str(s.get("body") or "").strip()
            bullets = s.get("bullets") or []

            if h:
                _add_paragraph(
                    document,
                    h.upper(),
                    bold=True,
                    size_pt=11,
                    space_before=8,
                    space_after=3,
                )
            if b:
                _add_paragraph(
                    document,
                    b,
                    bold=False,
                    size_pt=11,
                    space_before=0,
                    space_after=4,
                )
            if isinstance(bullets, list):
                for blt in bullets:
                    blt_str = str(blt or "").strip()
                    if not blt_str:
                        continue
                    if not blt_str.startswith("✅"):
                        blt_str = f"✅ {blt_str}"
                    _add_paragraph(
                        document,
                        blt_str,
                        bold=False,
                        size_pt=11,
                        space_before=0,
                        space_after=2,
                    )

    # 4. Mục khảo sát (Survey Section)
    survey_section = cp.get("survey_section")
    if isinstance(survey_section, dict) and (survey_section.get("text") or survey_section.get("feedback_summary")):
        is_training = (post_type == NewsPost.TYPE_FB_TRAINING or str(post_type).lower() in ("fb_training", "training"))
        default_heading = "📊 PHẢN HỒI SAU TẬP HUẤN" if is_training else "PHẢN HỒI SAU BUỔI LÀM VIỆC"
        s_head = str(survey_section.get("heading") or default_heading).strip()
        s_text = str(survey_section.get("text") or "").strip()
        s_fb = str(survey_section.get("feedback_summary") or "").strip()

        if s_head:
            _add_paragraph(
                document,
                s_head,
                bold=True,
                size_pt=11,
                space_before=8,
                space_after=3,
            )
        if s_text:
            _add_paragraph(
                document,
                s_text,
                bold=False,
                size_pt=11,
                space_before=0,
                space_after=4,
            )
        if s_fb:
            _add_paragraph(
                document,
                s_fb,
                bold=False,
                size_pt=11,
                space_before=0,
                space_after=6,
            )

    # 5. Lời cảm ơn & Kết (Closing)
    closing = str(cp.get("closing") or "").strip()
    if closing:
        _add_paragraph(
            document,
            closing,
            bold=False,
            size_pt=11,
            space_before=6,
            space_after=8,
        )

    # 6. Phần quảng bá (Chỉ áp dụng cho bài đăng Facebook/mạng xã hội, KHÔNG có trong Bài báo)
    if post_type != NewsPost.TYPE_PRESS_ARTICLE:
        slogan = str(cp.get("slogan") or "").strip()
        if slogan:
            _add_paragraph(
                document,
                slogan,
                bold=True,
                size_pt=11,
                space_before=6,
                space_after=4,
            )

        # Thông tin liên hệ: TÁCH DỌC, mỗi liên hệ 1 dòng riêng biệt
        contact_info = str(cp.get("contact_info") or "").strip()
        contact_lines = split_contact_info_lines(contact_info)
        for cl in contact_lines:
            _add_paragraph(
                document,
                cl,
                bold=False,
                size_pt=11,
                space_before=0,
                space_after=2,
            )

        # Hashtag: MỘT đoạn duy nhất, cách nhau đúng 1 dấu cách, căn TRÁI (không justify)
        hashtags = cp.get("hashtags") or []
        hashtag_str = format_hashtags_string(hashtags)
        if hashtag_str:
            _add_paragraph(
                document,
                hashtag_str,
                bold=False,
                size_pt=11,
                align=WD_ALIGN_PARAGRAPH.LEFT,
                space_before=6,
                space_after=6,
            )

    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    docx_buffer.seek(0)
    return docx_buffer
