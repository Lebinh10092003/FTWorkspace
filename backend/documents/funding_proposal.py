"""Builds the "Phiếu đề xuất kinh phí" as a .docx, ready to print and sign.

The layout follows the company's Google Docs master: a two-column letterhead,
numbered sections 1-5, a cost table that carries its own CỘNG / VAT / CHI PHÍ
KHÁC / TỔNG CỘNG rows, and a four-column signature strip. Word renders this
directly; nothing here depends on the browser that filled the form in.
"""
from datetime import date, datetime
from io import BytesIO

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

TICKED = "☒"      # ☒
UNTICKED = "☐"    # ☐
FONT = "Times New Roman"

PAYMENT_METHODS = [("transfer", "Chuyển khoản"), ("cash", "Tiền mặt")]
PROPOSAL_KINDS = [("new", "Mới"), ("adjustment", "Điều chỉnh, bổ sung")]
ATTACHMENTS = [
    ("quote", "Báo giá đủ thuế/phí"),
    ("quantity", "Bảng/danh sách chốt số lượng"),
    ("approval", "Phê duyệt trước"),
    ("spec", "Quy cách/mẫu sản phẩm"),
]
DECISIONS = [
    ("approved", "Đồng ý"),
    ("conditional", "Đồng ý có điều kiện"),
    ("more", "Yêu cầu bổ sung"),
    ("rejected", "Không đồng ý"),
]


def _money(value):
    """1234567 -> "1.234.567". Blank stays blank so the form can be printed empty."""
    if value in (None, "", False):
        return ""
    try:
        number = float(str(value).replace(",", "").replace(" ", ""))
    except (TypeError, ValueError):
        return str(value)
    if number == int(number):
        return f"{int(number):,}".replace(",", ".")
    return f"{number:,.2f}".replace(",", "\u0000").replace(".", ",").replace("\u0000", ".")


def _currency_unit(data):
    code = str(data.get("currency") or "VND").strip().upper()
    if len(code) != 3 or not code.isascii() or not code.isalpha():
        code = "VND"
    return "đồng" if code == "VND" else code


def _dotted(value, width=30):
    text = str(value or "").strip()
    return text if text else "…" * width


def _box(checked, label):
    return f"{TICKED if checked else UNTICKED} {label}"


def _parse_date(value):
    text = str(value or "").strip()
    if not text:
        return None
    for pattern in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _date_phrase(value, prefix="Hà Nội, ngày"):
    parsed = _parse_date(value)
    if not parsed:
        return f"{prefix} … tháng … năm …"
    return f"{prefix} {parsed.day:02d} tháng {parsed.month:02d} năm {parsed.year}"


def _short_date(value):
    parsed = _parse_date(value)
    return parsed.strftime("%d/%m/%Y") if parsed else "……/……/………"


def _style(run, *, bold=False, italic=False, size=13):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def _para(container, text="", *, bold=False, italic=False, size=13,
          align=None, space_after=2, space_before=0, indent=None):
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.space_before = Pt(space_before)
    paragraph.paragraph_format.line_spacing = 1.15
    if indent is not None:
        paragraph.paragraph_format.left_indent = Cm(indent)
    if align is not None:
        paragraph.alignment = align
    _style(paragraph.add_run(text), bold=bold, italic=italic, size=size)
    return paragraph


def _borderless(table):
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "none")
        borders.append(element)
    table._tbl.tblPr.append(borders)


def _cell_text(cell, text, *, bold=False, italic=False, size=13, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(1)
    paragraph.paragraph_format.space_before = Pt(1)
    _style(paragraph.add_run(text), bold=bold, italic=italic, size=size)


def _letterhead(document, data):
    table = document.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borderless(table)
    table.columns[0].width = Cm(6.5)
    table.columns[1].width = Cm(10.0)

    _cell_text(table.cell(0, 0), "CÔNG TY CỔ PHẦN", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(table.cell(0, 0), "CÔNG NGHỆ FERMAT", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _cell_text(table.cell(0, 1), "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _para(table.cell(0, 1), "Độc lập - Tự do - Hạnh phúc", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)

    number = _document_number(data)
    _cell_text(
        table.cell(1, 0),
        f"Số: {number}" if number else "Số: ……/PĐXKP-FT",
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _cell_text(
        table.cell(1, 1),
        _date_phrase(data.get("issuedOn")),
        italic=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    document.add_paragraph()


def _cost_table(document, data):
    items = [item for item in (data.get("items") or []) if isinstance(item, dict)]
    rows = max(len(items), 1)
    table = document.add_table(rows=rows + 5, cols=6)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(1.2), Cm(6.4), Cm(1.6), Cm(1.8), Cm(2.6), Cm(2.8)]
    headers = ["TT", "Nội dung, quy cách\nvà cấu phần chi phí", "ĐVT",
               "Số\nlượng", "Đơn giá\nchưa thuế", "Thành tiền\nchưa thuế"]
    for index, (label, width) in enumerate(zip(headers, widths)):
        table.columns[index].width = width
        _cell_text(table.cell(0, index), label, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)

    subtotal = 0.0
    for offset in range(rows):
        item = items[offset] if offset < len(items) else {}
        amount = _line_amount(item)
        subtotal += amount
        cells = [
            (str(offset + 1), WD_ALIGN_PARAGRAPH.CENTER),
            (_dotted(item.get("description"), 24), WD_ALIGN_PARAGRAPH.LEFT),
            (_dotted(item.get("unit"), 4), WD_ALIGN_PARAGRAPH.CENTER),
            (_dotted(item.get("quantity"), 4), WD_ALIGN_PARAGRAPH.CENTER),
            (_dotted(_money(item.get("unitPrice")), 6), WD_ALIGN_PARAGRAPH.RIGHT),
            (_dotted(_money(amount) if amount else "", 6), WD_ALIGN_PARAGRAPH.RIGHT),
        ]
        for index, (text, align) in enumerate(cells):
            _cell_text(table.cell(offset + 1, index), text, size=12, align=align)

    vat_rate = _number(data.get("vatRate"))
    vat_amount = subtotal * vat_rate / 100 if vat_rate else _number(data.get("vatAmount"))
    other = _number(data.get("otherCost"))
    total = subtotal + vat_amount + other

    summary = [
        ("CỘNG", _money(subtotal) if subtotal else ""),
        (f"Thuế GTGT VAT ({_trim(vat_rate)}%)" if vat_rate else "Thuế GTGT VAT ( …%)",
         _money(vat_amount) if vat_amount else ""),
        ("CHI PHÍ KHÁC", _money(other) if other else ""),
        ("TỔNG CỘNG", _money(total) if total else ""),
    ]
    for offset, (label, value) in enumerate(summary):
        row = rows + 1 + offset
        merged = table.cell(row, 0).merge(table.cell(row, 4))
        strong = label in {"CỘNG", "TỔNG CỘNG"}
        _cell_text(merged, label, bold=strong, size=12, align=WD_ALIGN_PARAGRAPH.RIGHT)
        _cell_text(table.cell(row, 5), _dotted(value, 6), bold=strong, size=12,
                   align=WD_ALIGN_PARAGRAPH.RIGHT)
    return total


def _number(value):
    try:
        return float(str(value).replace(",", "").replace(" ", "")) if value not in (None, "") else 0.0
    except (TypeError, ValueError):
        return 0.0


def _trim(value):
    return str(int(value)) if float(value) == int(value) else str(value)


def _line_amount(item):
    explicit = _number(item.get("amount"))
    if explicit:
        return explicit
    return _number(item.get("quantity")) * _number(item.get("unitPrice"))


def _signature_strip(document):
    table = document.add_table(rows=1, cols=4)
    _borderless(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = ["NGƯỜI LẬP\nPHIẾU", "PHỤ TRÁCH\nBỘ PHẬN", "KẾ TOÁN\nKIỂM TRA", "NGƯỜI\nPHÊ DUYỆT"]
    for index, label in enumerate(labels):
        table.columns[index].width = Cm(4.0)
        cell = table.cell(0, index)
        _cell_text(cell, label, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
        for _ in range(3):
            _para(cell, "", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        _para(cell, "…………………", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)


def build_funding_proposal(data):
    """Render the form and return (filename, docx bytes)."""
    data = data if isinstance(data, dict) else {}
    currency_unit = _currency_unit(data)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.6), Cm(1.6)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(13)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _letterhead(document, data)

    _para(document, "PHIẾU ĐỀ XUẤT KINH PHÍ", bold=True, size=16,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    _para(document, "Dùng cho đề xuất mới và điều chỉnh kinh phí", italic=True, size=12,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    _para(document, "Kính gửi: Tổng Giám đốc Công ty Cổ phần Công nghệ Fermat.", space_after=8)

    kind = str(data.get("kind") or "new")
    _para(document, "1. Thông tin đề xuất", bold=True, space_before=4)
    _para(document, "Loại đề xuất: "
          + "   ".join(_box(kind == value, label) for value, label in PROPOSAL_KINDS)
          + f"     Lần trình: {_dotted(data.get('submissionRound'), 4)}")
    _para(document, f"Người đề xuất: {_dotted(data.get('proposer'), 18)}"
                    f"     Bộ phận: {_dotted(data.get('department'), 16)}")
    _para(document, f"Công việc/dự án: {_dotted(data.get('project'), 40)}")
    _para(document, f"Trưởng bộ phận: {_dotted(data.get('departmentHead'), 40)}")
    _para(document, f"Mục đích, kết quả cần đạt: {_dotted(data.get('purpose'), 34)}")
    _para(document, f"Hạn cần duyệt: {_short_date(data.get('approvalDeadline'))}", space_after=8)

    _para(document, "2. Dự toán và phương án thực hiện", bold=True, space_before=4)
    _para(document, f"Đơn vị tiền: {currency_unit}.", italic=True, size=12)
    total = _cost_table(document, data)
    _para(document, f"Tổng kinh phí đề nghị (A + B + C): {_dotted(_money(total) if total else '', 20)} {currency_unit}.",
          space_before=6)
    _para(document, f"Thông tin nhà cung cấp: {_dotted(data.get('supplier'), 36)}")
    payment = str(data.get("paymentMethod") or "transfer")
    _para(document, "Thanh toán: "
          + "   ".join(_box(payment == value, label) for value, label in PAYMENT_METHODS) + ";")
    _para(document,
          f"Tạm ứng (nếu có): {_dotted(_money(data.get('advanceAmount')), 14)} {currency_unit} "
          f"tại phiếu số: {_dotted(data.get('advanceDocument'), 6)} "
          f"{_date_phrase(data.get('advanceDate'), prefix='ngày')}", space_after=8)

    _para(document, "3. Đối chiếu điều chỉnh và hồ sơ kèm theo", bold=True, space_before=4)
    _para(document, "Chỉ điền phần đối chiếu khi điều chỉnh; đề xuất mới ghi “Không áp dụng”.",
          italic=True, size=12)
    _para(document, f"Căn cứ lần duyệt trước (số phiếu/email, ngày, người duyệt): "
                    f"{_dotted(data.get('previousApproval'), 16)}")
    _comparison_table(document, data)
    _para(document, f"Lý do thay đổi; tác động đến tiến độ, ngân sách: {_dotted(data.get('changeReason'), 24)}",
          space_before=6)
    _para(document, f"Tổng sau điều chỉnh: {_dotted(_money(data.get('adjustedTotal')), 12)} {currency_unit}; "
                    f"đã chi/cam kết: {_dotted(_money(data.get('committedAmount')), 12)} {currency_unit}.")
    _para(document, f"Số tiền còn phải bố trí: {_dotted(_money(data.get('remainingAmount')), 12)} {currency_unit}; "
                    f"phần xin tăng: {_dotted(_money(data.get('increaseAmount')), 10)} {currency_unit}.")

    selected = set(data.get("attachments") or [])
    first = ATTACHMENTS[:2]
    rest = ATTACHMENTS[2:]
    _para(document, "Hồ sơ kèm theo: "
          + "   ".join(_box(value in selected, label) for value, label in first))
    other_label = str(data.get("attachmentOther") or "").strip()
    _para(document, "   ".join(_box(value in selected, label) for value, label in rest)
          + "   " + _box(bool(other_label), f"Tài liệu khác: {_dotted(other_label, 10)}"))
    _para(document, f"Đường dẫn hồ sơ BNDC: {_dotted(data.get('bndcLink'), 30)}", space_after=8)

    _para(document, "4. Xác nhận và ý kiến kiểm tra", bold=True, space_before=4)
    _para(document, "Người lập chịu trách nhiệm về nhu cầu, số lượng, cấu phần chi phí và việc không "
                    "đề xuất trùng khoản đã được duyệt hoặc thanh toán.", space_after=8)

    decision = str(data.get("decision") or "")
    _para(document, "5. Phê duyệt của người có thẩm quyền", bold=True, space_before=4)
    _para(document, "   ".join(_box(decision == value, label) for value, label in DECISIONS))
    _para(document, f"Tổng mức được duyệt (gồm thuế, phí): {_dotted(_money(data.get('approvedTotal')), 20)} {currency_unit}.")
    _para(document, f"Điều kiện; người thực hiện; thời hạn: {_dotted(data.get('conditions'), 24)}")
    _para(document, "Ký, ghi rõ họ tên, ngày ký; người phê duyệt ghi thêm chức vụ.",
          italic=True, size=12, space_after=6)

    _signature_strip(document)
    _para(document, "Phiếu dùng để duyệt kinh phí; hồ sơ tạm ứng, thanh toán thực hiện riêng. "
                    "Phát sinh vượt mức/phạm vi đã duyệt phải trình lại trước khi cam kết chi.",
          italic=True, size=11, space_before=8)

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return _filename(data), stream.getvalue()


def _comparison_table(document, data):
    comparison = data.get("comparison") if isinstance(data.get("comparison"), dict) else {}
    table = document.add_table(rows=3, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Nội dung", "Đã duyệt trước", "Đề nghị lần này", "Tăng/giảm"]
    widths = [Cm(5.0), Cm(3.6), Cm(3.6), Cm(3.4)]
    for index, (label, width) in enumerate(zip(headers, widths)):
        table.columns[index].width = width
        _cell_text(table.cell(0, index), label, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    labels = [
        ("quantity", "Số lượng/đơn giá\nhạng mục thay đổi"),
        ("total", "Tổng kinh phí\ncùng phạm vi và cơ sở thuế"),
    ]
    for offset, (key, label) in enumerate(labels):
        entry = comparison.get(key) if isinstance(comparison.get(key), dict) else {}
        _cell_text(table.cell(offset + 1, 0), label, size=12)
        for column, field in enumerate(("previous", "current", "delta"), start=1):
            _cell_text(table.cell(offset + 1, column), _dotted(entry.get(field), 8), size=12,
                       align=WD_ALIGN_PARAGRAPH.CENTER)


def _filename(data):
    number = _document_number(data).replace("/", "-")
    stamp = _parse_date(data.get("issuedOn")) or date.today()
    tail = number or stamp.strftime("%Y%m%d")
    return f"Phieu-de-xuat-kinh-phi-{tail}.docx"


def _document_number(data):
    prefix = str(data.get("documentNumber") or "").split("/", 1)[0]
    digits = "".join(char for char in prefix if char.isascii() and char.isdigit())
    return f"{digits}/PĐXKP-FT" if digits else ""
